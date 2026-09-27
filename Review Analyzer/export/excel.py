from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PCT_COLUMNS = {"Review Rate", "Overall Coverage", "Good Rate", "Bad Rate"}


def _style_sheet(ws) -> None:
    header_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        header = str(col_cells[0].value or "")

        if header in PCT_COLUMNS:
            for cell in col_cells[1:]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.0%"

        if header == "Review Text":
            ws.column_dimensions[letter].width = 80
            for cell in col_cells:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        else:
            max_len = max(
                (len(str(c.value)) if c.value is not None else 0 for c in col_cells),
                default=8,
            )
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 38)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _build_status_df(collection_status: dict) -> pd.DataFrame:
    rows = []
    for ch, info in collection_status.items():
        profiles = info.get("profiles") or {}
        if profiles:
            for profile, pinfo in profiles.items():
                rows.append({
                    "Agency": ch,
                    "Profile": profile,
                    "Area": ", ".join(pinfo.get("areas") or []),
                    "Status": pinfo.get("status", ""),
                    "Collected Reviews": pinfo.get("count", 0),
                    "Error": pinfo.get("error", ""),
                })
        else:
            rows.append({
                "Agency": ch,
                "Profile": "",
                "Area": "",
                "Status": info.get("status", ""),
                "Collected Reviews": info.get("count", 0),
                "Error": info.get("error", ""),
            })
    df = pd.DataFrame(rows)
    # 채널 단위로 '수집 → 매칭' 을 함께 보여준다. 수집 개수만 적으면 OTA 화면의
    # 개수와 보고서 숫자가 왜 다른지 설명할 수 없다.
    if not df.empty:
        df["Matched to Booking"] = [
            collection_status.get(ch, {}).get("matched", "") for ch in df["Agency"]
        ]
        df["Outside Scope"] = [
            collection_status.get(ch, {}).get("unmatched", "") for ch in df["Agency"]
        ]
        df["Note"] = [
            "Collected = 이 계정에서 가져온 리뷰(예약코드 기준) · "
            "Matched = 이번 분석 범위 예약과 연결된 수 · "
            "Outside Scope = 범위 밖 예약의 리뷰"
            if str(st).startswith(("SUCCESS", "PARTIAL")) else ""
            for st in df["Status"]
        ]
    return df


def _save_workbook(
    out_path: Path,
    master: pd.DataFrame,
    overall: pd.DataFrame,
    tours: pd.DataFrame,
    areas: pd.DataFrame,
    channels: pd.DataFrame,
    unsupported: pd.DataFrame,
    status_df: pd.DataFrame,
) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        overall.to_excel(writer, sheet_name="Overall", index=False)
        tours.to_excel(writer, sheet_name="Tour Metrics", index=False)
        areas.to_excel(writer, sheet_name="Area Metrics", index=False)
        channels.to_excel(writer, sheet_name="Channel Metrics", index=False)
        unsupported.to_excel(writer, sheet_name="Unsupported", index=False)
        status_df.to_excel(writer, sheet_name="Collection Status", index=False)
        master.to_excel(writer, sheet_name="Master", index=False)

        for ws in writer.book.worksheets:
            _style_sheet(ws)


def save_phase1_excel(
    path: str,
    master: pd.DataFrame,
    overall: pd.DataFrame,
    tours: pd.DataFrame,
    areas: pd.DataFrame,
    channels: pd.DataFrame,
    unsupported: pd.DataFrame,
    collection_status: dict,
) -> str:
    out_path = Path(path)
    status_df = _build_status_df(collection_status)

    try:
        _save_workbook(
            out_path,
            master,
            overall,
            tours,
            areas,
            channels,
            unsupported,
            status_df,
        )
        return str(out_path)
    except PermissionError:
        # Most commonly the previous result workbook is still open in Excel.
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        alt_path = out_path.with_name(f"{out_path.stem}_{stamp}{out_path.suffix}")
        try:
            _save_workbook(
                alt_path,
                master,
                overall,
                tours,
                areas,
                channels,
                unsupported,
                status_df,
            )
            return str(alt_path)
        except PermissionError as exc:
            raise RuntimeError(
                "결과 Excel을 저장할 수 없습니다. 같은 결과 파일이 Excel에서 열려 있다면 닫은 후 다시 실행해 주세요."
            ) from exc
